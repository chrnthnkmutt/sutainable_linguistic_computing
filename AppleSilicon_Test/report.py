"""Excel report writer.

Workbook sheets
  Summary      one row per measured run (appended on every execution)
  Samples      time series (CPU %, GPU %, power, RAM, VRAM) for charting
  Statistics   mean / std-dev across the runs of the latest session
  Environment  versions, detected hardware, backends, warnings, metric definitions
"""
from __future__ import annotations

import statistics
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NA = "N/A"

SUMMARY_HEADERS = [
    "Timestamp", "Session ID", "Run #", "Model alias", "Prompt", "Selected hardware",
    "Thinking level", "Thinking method",
    "Inference duration (s)", "Time to first token (s)",
    "Tokens generated (~stream chunks)", "Tokens/sec",
    "Process RAM start (MB)", "Process RAM end (MB)", "Process RAM peak (MB)",
    "System RAM % (start -> end)", "System RAM used peak (MB)",
    "GPU VRAM start (MB)", "GPU VRAM end (MB)", "GPU VRAM peak (MB)",
    "CPU hardware", "CPU usage change (% -> %)", "CPU usage avg (%)", "CPU usage peak (%)",
    "CPU power avg (W)", "CPU power peak (W)", "CPU energy (J)", "CPU net energy vs idle (J)",
    "GPU hardware", "GPU usage change (% -> %)", "GPU usage avg (%)", "GPU usage peak (%)",
    "GPU power avg (W)", "GPU power peak (W)", "GPU energy (J)", "GPU net energy vs idle (J)",
    "DRAM energy (J)", "GPU SRAM energy (J)", "Apple Neural Engine energy (J)",
    "Total energy (J)", "Total energy (Wh)", "Avg total power (W)", "Net total energy vs idle (J)",
    "Energy per token (J/token)", "Notes",
]
SAMPLE_HEADERS = ["Session ID", "Run #", "Source", "Time (s)", "Usage (%)", "Power (W)",
                  "Process RAM (MB)", "System RAM (%)", "VRAM (MB)"]
ENV_HEADERS = ["Session ID", "Timestamp", "Item", "Value"]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _v(x: Any, nd: int = 3) -> Any:
    if x is None:
        return NA
    if isinstance(x, float):
        return round(x, nd)
    return x


def _change(a: Optional[float], b: Optional[float]) -> str:
    return NA if a is None or b is None else f"{a:.1f} -> {b:.1f}"


def _component(res: Optional[Dict[str, Any]], idle: Dict[str, Any], prefix: str) -> Dict[str, Any]:
    if res is None:
        return {f"{prefix} hardware": "Not selected", **{h: NA for h in SUMMARY_HEADERS
                if h.startswith(prefix + " ") and h != f"{prefix} hardware"}}
    net = None
    ip = idle.get(res["key"], {}).get("power_w")
    if res["energy_j"] is not None and ip is not None:
        net = max(0.0, res["energy_j"] - ip * res["duration_s"])
    return {
        f"{prefix} hardware": res["name"],
        f"{prefix} usage change (% -> %)": _change(res["usage_start"], res["usage_end"]),
        f"{prefix} usage avg (%)": _v(res["usage_avg"], 1),
        f"{prefix} usage peak (%)": _v(res["usage_peak"], 1),
        f"{prefix} power avg (W)": _v(res["power_avg_w"]),
        f"{prefix} power peak (W)": _v(res["power_peak_w"]),
        f"{prefix} energy (J)": _v(res["energy_j"]),
        f"{prefix} net energy vs idle (J)": _v(net),
    }


def build_row(session: Dict[str, Any], run: Dict[str, Any]) -> Dict[str, Any]:
    res, mem, idle = run["results"], run["memory"], session["idle"]
    cpu = next((res[m["key"]] for m in session["selected"] if m["kind"] == "cpu"), None)
    gpu = next((res[m["key"]] for m in session["selected"] if m["kind"] == "gpu"), None)
    cpu_row, gpu_row = _component(cpu, idle, "CPU"), _component(gpu, idle, "GPU")

    extra: Dict[str, Any] = {}
    for r in res.values():
        extra.update(r.get("extra", {}))

    energies = [r["energy_j"] for r in (cpu, gpu) if r is not None and r["energy_j"] is not None]
    nets = [v for v in (cpu_row["CPU net energy vs idle (J)"], gpu_row["GPU net energy vs idle (J)"])
            if v != NA]
    total_j = sum(energies) if energies else None
    dur = run["duration_s"]
    tokens = run["tokens"]

    notes: List[str] = []
    for r in list(res.values()) + [mem, {"name": "Run", "notes": run.get("notes", [])}]:
        for n in r.get("notes", []):
            tag = f"[{r['name']}] {n}"
            if tag not in notes:
                notes.append(tag)
    selected_comp = [r for r in (cpu, gpu) if r is not None]
    if total_j is not None and len(energies) < len(selected_comp):
        notes.append("Total covers only the components with a measured energy value")

    nvidia = res.get("nvidia_gpu", {}).get("extra", {})
    unified = "apple_gpu" in res or "apple_cpu" in res

    def vram(k: str) -> Any:
        if k in nvidia and nvidia[k] is not None:
            return round(nvidia[k], 1)
        return "N/A (unified memory)" if unified else NA

    row: Dict[str, Any] = {
        "Timestamp": run["timestamp"], "Session ID": session["id"], "Run #": run["run"],
        "Model alias": session["alias"], "Prompt": session["prompt"],
        "Selected hardware": ", ".join(m["label"] for m in session["selected"]) or "None",
        "Thinking level": session.get("thinking", NA),
        "Thinking method": session.get("thinking_method", NA),
        "Inference duration (s)": _v(dur),
        "Time to first token (s)": _v(run["ttft_s"]),
        "Tokens generated (~stream chunks)": tokens,
        "Tokens/sec": _v(tokens / dur if dur > 0 else None, 2),
        "Process RAM start (MB)": _v(mem["rss_start_mb"], 1),
        "Process RAM end (MB)": _v(mem["rss_end_mb"], 1),
        "Process RAM peak (MB)": _v(mem["rss_peak_mb"], 1),
        "System RAM % (start -> end)": _change(mem["sys_pct_start"], mem["sys_pct_end"]),
        "System RAM used peak (MB)": _v(mem["sys_used_peak_mb"], 1),
        "GPU VRAM start (MB)": vram("vram_start_mb"),
        "GPU VRAM end (MB)": vram("vram_end_mb"),
        "GPU VRAM peak (MB)": vram("vram_peak_mb"),
        **cpu_row, **gpu_row,
        "DRAM energy (J)": _v(extra.get("dram_j")),
        "GPU SRAM energy (J)": _v(extra.get("gpu_sram_j")),
        "Apple Neural Engine energy (J)": _v(extra.get("ane_j")),
        "Total energy (J)": _v(total_j),
        "Total energy (Wh)": _v(total_j / 3600 if total_j is not None else None, 6),
        "Avg total power (W)": _v(total_j / dur if total_j is not None and dur > 0 else None),
        "Net total energy vs idle (J)": _v(sum(nets) if nets else None),
        "Energy per token (J/token)": _v(total_j / tokens if total_j is not None and tokens else None, 4),
        "Notes": "; ".join(notes),
    }
    return {h: row.get(h, NA) for h in SUMMARY_HEADERS}


def _ensure_sheet(wb: Workbook, title: str, headers: List[str], widths: Optional[Dict[str, int]] = None):
    name, n = title, 1
    while name in wb.sheetnames:
        ws = wb[name]
        if [c.value for c in ws[1]][: len(headers)] == headers:
            return ws
        n += 1
        name = f"{title} (v{n})"  # layout changed since the file was created
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, min(len(h) + 2, 26)))
    return ws


def _statistics_sheet(wb: Workbook, rows: List[Dict[str, Any]], session_id: str) -> None:
    if "Statistics" in wb.sheetnames:
        del wb["Statistics"]
    ws = wb.create_sheet("Statistics")
    ws.append([f"Latest session: {session_id}  ({len(rows)} measured run(s))"])
    ws["A1"].font = Font(bold=True)
    ws.append(["Metric", "Mean", "Std dev (sample)", "N"])
    for c in ws[2]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for h in SUMMARY_HEADERS:
        vals = [r[h] for r in rows]
        nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
        if h in ("Run #",) or not nums or len(nums) != len(vals):
            continue
        ws.append([h, round(statistics.fmean(nums), 4),
                   round(statistics.stdev(nums), 4) if len(nums) > 1 else NA, len(nums)])
    ws.column_dimensions["A"].width = 38
    for col in "BCD":
        ws.column_dimensions[col].width = 18


def _save(wb: Workbook, path: Path) -> Path:
    try:
        wb.save(path)
        return path
    except PermissionError:  # e.g. file open in Excel
        alt = path.with_name(f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}")
        wb.save(alt)
        return alt


def write_report(path: Path, session: Dict[str, Any], runs: List[Dict[str, Any]],
                 env_items: List[tuple]) -> Path:
    path = Path(path)
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    summary = _ensure_sheet(wb, "Summary", SUMMARY_HEADERS,
                            {"Timestamp": 20, "Session ID": 12, "Prompt": 34, "Model alias": 28,
                             "Selected hardware": 30, "Thinking method": 40, "CPU hardware": 34, "GPU hardware": 34,
                             "Notes": 60})
    samples = _ensure_sheet(wb, "Samples", SAMPLE_HEADERS)
    env = _ensure_sheet(wb, "Environment", ENV_HEADERS, {"Item": 34, "Value": 90, "Timestamp": 20})

    rows = []
    for run in runs:
        row = build_row(session, run)
        rows.append(row)
        summary.append([row[h] for h in SUMMARY_HEADERS])
        for key, r in run["results"].items():
            for s in r["samples"]:
                samples.append([session["id"], run["run"], r["name"], s["t"], _v(s.get("usage"), 1),
                                _v(s.get("power")), None, None, _v(s.get("vram_mb"), 1) if "vram_mb" in s else None])
        for s in run["memory"]["samples"]:
            samples.append([session["id"], run["run"], "Memory", s["t"], None, None,
                            round(s["rss_mb"], 1), s["sys_pct"], None])

    for item, value in env_items:
        env.append([session["id"], session["started"], item, str(value)])

    if rows:
        _statistics_sheet(wb, rows, session["id"])
    wb.move_sheet("Summary", offset=-wb.index(wb["Summary"]))
    return _save(wb, path)
