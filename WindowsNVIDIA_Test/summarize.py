"""Summarize Sustain_SLM xlsx reports into one table per condition (one xlsx = one condition).

    python summarize.py results/*.xlsx            -> results/summary.md, summary_by_condition.csv,
                                                     summary_by_task.csv

Accuracy uses prompts/gold.json (QA: any gold string in the answer; sentiment: earliest label
mentioned). Summaries and jokes are not auto-scored.
"""
from __future__ import annotations

import csv
import json
import statistics
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).parent
GOLD = json.loads((HERE / "prompts" / "gold.json").read_text(encoding="utf-8"))


def rows_of(path: Path):
    wb = load_workbook(path, read_only=True)
    name = [n for n in wb.sheetnames if n.startswith("Summary")][-1]  # newest layout
    it = wb[name].iter_rows(values_only=True)
    head = next(it)
    return [dict(zip(head, r)) for r in it if any(v is not None for v in r)]


def num(x):
    return x if isinstance(x, (int, float)) and not isinstance(x, bool) else None


def score(pid: str, text: str):
    g = GOLD["items"].get(pid)
    if not g:
        return None
    t = str(text or "").lower()  # no answer (e.g. cut by the context limit) counts as wrong
    if g["task"] == "qa":
        return float(any(s.lower() in t for s in g["gold_any"]))
    if g["task"] == "classify":
        labels = {**GOLD["labels"]["th"], **GOLD["labels"]["en"]}
        hits = [(t.find(k.lower()), v) for k, v in labels.items() if k.lower() in t]
        return float(min(hits)[1] == g["gold"]) if hits else 0.0
    return None


def mean(xs):
    xs = [x for x in xs if x is not None]
    return statistics.fmean(xs) if xs else None


def ci95(xs):
    xs = [x for x in xs if x is not None]
    return 1.96 * statistics.stdev(xs) / len(xs) ** 0.5 if len(xs) > 1 else None


def agg(rows):
    col = lambda k: [num(r.get(k)) for r in rows]
    tot = col("Total energy (J)")
    acc = [score(str(r.get("Prompt ID")), r.get("Response")) for r in rows]
    toks = sum(x for x in col("Tokens generated (~stream chunks)") if x)
    return {
        "n": len(rows),
        "cpu_j": mean(col("CPU energy (J)")), "gpu_j": mean(col("GPU energy (J)")),
        "total_j": mean(tot), "total_ci95": ci95(tot),
        "net_cpu_j": mean(col("CPU net energy vs idle (J)")), "net_gpu_j": mean(col("GPU net energy vs idle (J)")),
        "net_total_j": mean(col("Net total energy vs idle (J)")),
        "j_per_token": sum(x for x in tot if x) / toks if toks else None,
        "tokens": mean(col("Tokens generated (~stream chunks)")),
        "duration_s": mean(col("Inference duration (s)")),
        "ttft_s": statistics.median([x for x in col("Time to first token (s)") if x is not None] or [0]),
        "gen_tps": mean(col("Generation tokens/sec")),
        "cpu_w": mean(col("CPU power avg (W)")), "gpu_w": mean(col("GPU power avg (W)")),
        "ram_peak_mb": max([x for x in col("Process RAM peak (MB)") if x] or [0]),
        "vram_peak_mb": max([x for x in col("GPU VRAM peak (MB)") if x] or [0]),
        "finished": mean([1.0 if r.get("Finish reason") == "stop" else 0.0 for r in rows]),
        "accuracy": mean(acc), "n_scored": sum(a is not None for a in acc),
    }


def fmt(v, nd=1):
    return "–" if v is None else (f"{v:.{nd}f}" if isinstance(v, float) else str(v))


def main(paths):
    paths = [Path(p) for p in paths if not Path(p).name.startswith("~$")]
    out_dir = paths[0].parent
    by_cond, by_task = {}, {}
    for p in paths:
        rows = rows_of(p)
        if not rows:
            continue
        cond = p.stem
        by_cond[cond] = agg(rows)
        groups = defaultdict(list)
        for r in rows:
            groups[GOLD["items"].get(str(r.get("Prompt ID")), {}).get("task", "other")].append(r)
        for task, rs in groups.items():
            by_task[(cond, task)] = agg(rs)

    keys = list(next(iter(by_cond.values())).keys())
    with open(out_dir / "summary_by_condition.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["condition"] + keys)
        for c, a in by_cond.items():
            w.writerow([c] + [a[k] for k in keys])
    with open(out_dir / "summary_by_task.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["condition", "task"] + keys)
        for (c, t), a in by_task.items():
            w.writerow([c, t] + [a[k] for k in keys])

    md = ["# Qwen3.5-4B on Ollama: Sustain_SLM measurements (CPU via Windows EMI + GPU via NVML)", "",
          "| Condition | n | CPU J | GPU J | Total J (±CI95) | Net total J | J/token | Tokens | Gen tok/s | TTFT s | CPU W | GPU W | RAM MB | VRAM MB | Finished | Accuracy |",
          "|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|"]
    for c, a in by_cond.items():
        md.append(f"| {c} | {a['n']} | {fmt(a['cpu_j'])} | {fmt(a['gpu_j'])} | {fmt(a['total_j'])} ± {fmt(a['total_ci95'])} | "
                  f"{fmt(a['net_total_j'])} | {fmt(a['j_per_token'], 2)} | {fmt(a['tokens'])} | {fmt(a['gen_tps'])} | "
                  f"{fmt(a['ttft_s'], 2)} | {fmt(a['cpu_w'])} | {fmt(a['gpu_w'])} | {fmt(a['ram_peak_mb'], 0)} | "
                  f"{fmt(a['vram_peak_mb'], 0)} | {fmt(a['finished'] * 100, 0)}% | "
                  f"{fmt(a['accuracy'] * 100 if a['accuracy'] is not None else None, 0)}% (n={a['n_scored']}) |")
    md += ["", "## By task", "", "| Condition | Task | n | Total J | Net total J | Tokens | Duration s | Finished | Accuracy |",
           "|---|---|---|---|---|---|---|---|---|"]
    for (c, t), a in by_task.items():
        md.append(f"| {c} | {t} | {a['n']} | {fmt(a['total_j'])} | {fmt(a['net_total_j'])} | {fmt(a['tokens'])} | "
                  f"{fmt(a['duration_s'], 2)} | {fmt(a['finished'] * 100, 0)}% | "
                  f"{fmt(a['accuracy'] * 100 if a['accuracy'] is not None else None, 0)}% |")
    (out_dir / "summary.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    print("\n".join(md))


if __name__ == "__main__":
    main(sys.argv[1:] or sorted(str(p) for p in (HERE / "results").glob("*.xlsx")))
